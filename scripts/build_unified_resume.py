from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"C:\Users\hp\Desktop\selfwebsite")
OUT = ROOT / "output" / "resumes"
OUT.mkdir(parents=True, exist_ok=True)
FILE = OUT / "Chelsea_Zhao_Product_Manager_Resume_CN.docx"
FONT_CN, FONT_EN = "Microsoft YaHei", "Aptos"
INK, MUTED, BLUE, GREEN, PALE, LINE = "172632", "64747E", "286B8F", "78A85C", "EBF6FA", "D9E8E5"

def font(run, size=8.65, bold=False, color=INK, italic=False):
    run.font.name = FONT_EN
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CN)
    run.font.size = Pt(size); run.bold = bold; run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)

def shade(p, fill):
    pPr = p._p.get_or_add_pPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill); shd.set(qn("w:val"), "clear"); pPr.append(shd)

def bottom_border(p):
    pPr=p._p.get_or_add_pPr(); bdr=OxmlElement("w:pBdr"); bottom=OxmlElement("w:bottom")
    for k,v in (("val","single"),("sz","5"),("space","2"),("color",LINE)): bottom.set(qn(f"w:{k}"),v)
    bdr.append(bottom); pPr.append(bdr)

def hyperlink(p, text, url):
    rid=p.part.relate_to(url,"http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",is_external=True)
    h=OxmlElement("w:hyperlink"); h.set(qn("r:id"),rid); r=OxmlElement("w:r"); pr=OxmlElement("w:rPr")
    c=OxmlElement("w:color"); c.set(qn("w:val"),BLUE); pr.append(c); sz=OxmlElement("w:sz"); sz.set(qn("w:val"),"17"); pr.append(sz)
    r.append(pr); t=OxmlElement("w:t"); t.text=text; r.append(t); h.append(r); p._p.append(h)

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.top_margin=Cm(.9); sec.bottom_margin=Cm(.8); sec.left_margin=Cm(1.25); sec.right_margin=Cm(1.25)
normal=doc.styles["Normal"]; normal.font.name=FONT_CN; normal._element.rPr.rFonts.set(qn("w:eastAsia"),FONT_CN); normal.font.size=Pt(8.65)
normal.paragraph_format.space_after=Pt(1.2); normal.paragraph_format.line_spacing=1.02
bullet=doc.styles["List Bullet"]; bullet.paragraph_format.left_indent=Cm(.42); bullet.paragraph_format.first_line_indent=Cm(-.22); bullet.paragraph_format.space_after=Pt(.55); bullet.paragraph_format.line_spacing=1.0

p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1); shade(p,PALE)
r=p.add_run("赵露 · CHELSEA ZHAO\n"); font(r,20,bold=True,color=BLUE)
r=p.add_run("产品经理  |  用户洞察 · AI 快速构建 · 数据验证 · 游戏化体验"); font(r,9.5,bold=True)
c=doc.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after=Pt(2)
for i,text in enumerate(["+86 150 3759 6216","chelsea299@163.com","吉隆坡，马来西亚"]):
    if i: font(c.add_run("  ·  "),8,color=GREEN)
    font(c.add_run(text),8)
font(c.add_run("  ·  "),8,color=GREEN); hyperlink(c,"rescueducks.xyz","https://www.rescueducks.xyz")

def heading(title,kicker):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(2.5); p.paragraph_format.space_after=Pt(1); bottom_border(p)
    font(p.add_run(title),10.6,bold=True,color=BLUE); font(p.add_run("   "+kicker),7,bold=True,color=GREEN)

def entry(title,meta,bullets,note=None):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(.6); p.paragraph_format.space_after=Pt(.2); p.paragraph_format.tab_stops.add_tab_stop(Cm(17.6),WD_TAB_ALIGNMENT.RIGHT)
    font(p.add_run(title),9,bold=True); font(p.add_run("\t"+meta),7.7,bold=True,color=MUTED)
    if note:
        n=doc.add_paragraph(); n.paragraph_format.space_after=Pt(.2); font(n.add_run(note),7.8,color=MUTED,italic=True)
    for b in bullets:
        q=doc.add_paragraph(style="List Bullet"); font(q.add_run(b),8.25)

heading("个人简介","PROFILE")
p=doc.add_paragraph(); font(p.add_run("关注用户体验、数据验证与 AI 落地的产品经理候选人。拥有教育一线、数据分析与互联网协作经验，能够从真实场景发现问题，定义 MVP，借助 AI 编程快速构建，并通过用户测试和行为证据持续迭代。"),8.35)

heading("产品项目","SELECTED PRODUCT WORK")
entry("Rescue Ducks｜雅思同义替换学习游戏","0→1 · 已上线",[
"识别“学生孤立背词、做题时无法识别同义替换”的真实问题，聚焦雅思阅读 7 分以下用户；提出以短回合游戏强化词间关系的产品假设。",
"定义并用 AI 编程实现 MVP：开始界面、同义词组匹配、双击释义、灯塔救援循环与复习界面；主动暂缓登录、排行榜及重特效，控制验证边界。",
"组织 30–40 名学生试玩；针对引导不清、双击能力不可发现和低基础用户易受挫，补充 onboarding、同组同色视觉支架与复习流程。",
])
entry("Agora｜家庭学习场景苏格拉底式数学 Agent","0→1 · Currently Building",[
"围绕通用 AI 容易直接泄露答案的问题，设计六级提示机制，并将概念误解、运算失误、步骤遗漏与表达不清纳入错误识别。",
"设计答案泄露控制、学习证据与智能练习路径；以正确性、提示适切性、答案泄露、错误识别、学习推进构建 AI 输出评测框架。",
])
entry("公益摄影服务平台","团队项目",["参与 15 人团队的用户招募、内容传播与服务交付，累计曝光 6,000+、服务 248+ 人次；项目获全国铜奖（前 1%），体现运营协同与服务设计能力。"])

heading("工作经历","EXPERIENCE")
entry("Kalowave｜数据分析实习生","2026.04 — 2026.06",[
"使用 Athena SQL 与 Metabase 分析 GMV、收入、佣金、转化率及用户分层，将业务问题转化为可追踪指标，支持经营判断。",
])
entry("新东方｜国际课程教师","2024.11 — 2025.09",[
"负责雅思阅读、IG 数学与 AP 微积分，基于学习卡点设计本地化内容和反馈路径；曾帮助 1 名 IG 学生在 10 次课后由 U 提升至 A。",
"将复杂概念拆解为分层提示和练习，积累教育用户需求、学习动机及效果评估的一线理解。",
])
entry("小红书｜商业／电商招聘实习生","2024.01 — 2024.04",[
"在高节奏互联网团队支持 10+ 岗位、100+ 面试，跟踪招聘漏斗并协同业务与候选人，适应快速变化与跨角色沟通。",
])

heading("教育背景","EDUCATION")
entry("马来亚大学｜应用统计学 硕士","2025.09 — 2027.07",["GPA 3.68 / 4.00；统计建模、数据分析与实验思维。"])
entry("兰州大学｜本科","2020.09 — 2024.07",["GPA 3.82 / 5.00；581 小时志愿服务经历。"])

heading("技能","SKILLS")
for label,value in [
    ("产品", "用户访谈、需求分析、MVP 定义、用户旅程、游戏化机制、A/B Test、AI 产品评测"),
    ("数据", "SQL、Athena、Metabase、Excel、统计分析、漏斗与用户分层"),
    ("构建", "AI 编程工具、快速原型、前端协作、Prompt 设计与模型输出验证"),
    ("语言", "中文（母语）、英文（工作语言）"),
]:
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(.55); font(p.add_run(label+"｜"),8,bold=True,color=BLUE); font(p.add_run(value),8)

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(footer.add_run("Chelsea Zhao · Turning real problems into products people want to use."),7,color=MUTED,italic=True)
doc.save(FILE)
print(FILE)
